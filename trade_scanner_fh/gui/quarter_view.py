"""
Quarter view (v8.1.0) — one ticker's quarterly earnings, transposed.

Double-clicking a ticker in the results table opens a window with the
quarters as COLUMNS and the per-quarter metrics as ROWS — the same values as
the Q-X blocks, laid out to read across time. The user's decisions
(2026-10-04) it follows:

  * columns: every quarter the scan produced for that ticker, whatever is
    hidden in the table; each header carries the quarter, its report date and
    fiscal quarter (so the two Q-X Date rows are not repeated as metrics);
  * rows: the per-quarter metrics in the scan that are not hidden — a metric
    is dropped when its Hide Q Columns type is ticked, or when EVERY one of
    its quarters is hidden individually;
  * a ✓ row per active Beats / Growth / Accelerating filter, under the
    quarters its run counted;
  * colours are the results table's own, cell for cell (a cell hidden in the
    table was never coloured there, so it shows uncoloured here);
  * a scan with no Q-X blocks shows Q-1 from the Current rows, or says how
    to get quarters;
  * a SNAPSHOT: later scans, hides or colour changes do not touch an open
    view; one window per double-click;
  * Copy (tab-separated, pastes into Excel) and Export to Excel (one sheet,
    real numbers, colours optional).

`build_quarter_view` is pure — no Qt — so the layout rules are tested
without a window. `QuarterViewDialog` renders a view; `write_quarter_view_
xlsx` writes one.
"""

from __future__ import annotations

import logging
import math
from dataclasses import dataclass, field
from datetime import date
from typing import Callable, Optional

import numpy as np
import pandas as pd
from PyQt6.QtCore import Qt
from PyQt6.QtGui import QColor
from PyQt6.QtWidgets import (
    QApplication, QCheckBox, QDialog, QFileDialog, QHBoxLayout, QLabel,
    QMessageBox, QPushButton, QTableWidget, QTableWidgetItem, QVBoxLayout,
)

from . import coloring as C
from .widgets import _Q_COL_RE, cell_tooltip

log = logging.getLogger("scanner.gui")

# (block suffix, row label, display format, Excel number format) per side, in
# the order the table's Q-X blocks show them. The display formats are the
# table's own (widgets._build_dynamic_columns) so a cell reads the same in
# both places; the Excel formats keep the numbers numeric.
_PCT_XL = '+0.00"%";-0.00"%";0.00"%"'
METRICS: dict = {
    "eps": (
        ("reported_eps", "Reported EPS", "{:.2f}", "0.00"),
        ("surprise_eps_dollar", "Surp EPS $", "{:+.2f}", "+0.00;-0.00;0.00"),
        ("surprise_eps_pct", "Surp EPS %", "{:+.2f}%", _PCT_XL),
        ("yoy_eps_pct", "YoY EPS %", "{:+.2f}%", _PCT_XL),
    ),
    "rev": (
        ("reported_rev", "Reported Rev", "{:,.1f}", "#,##0.0"),
        ("surprise_rev_dollar", "Surp Rev $", "{:+,.1f}",
         "+#,##0.0;-#,##0.0;0.0"),
        ("surprise_rev_pct", "Surp Rev %", "{:+.2f}%", _PCT_XL),
        ("yoy_rev_pct", "YoY Rev %", "{:+.2f}%", _PCT_XL),
    ),
}
# The Current (most recent quarter) columns, for a scan without Q-X blocks.
# Same suffixes as the blocks, unprefixed.
_CURRENT_KEYS = tuple(s for side in ("eps", "rev") for s, *_ in METRICS[side])
MARK = "✓"


@dataclass
class QuarterHead:
    k: int
    report_date: Optional[pd.Timestamp]
    fiscal: Optional[pd.Timestamp]

    def label_lines(self) -> list:
        lines = [f"Q-{self.k}"]
        if self.report_date is not None:
            lines.append(self.report_date.strftime("%Y-%m-%d"))
        if self.fiscal is not None:
            lines.append("FQ " + self.fiscal.strftime("%Y-%m"))
        return lines


@dataclass
class Cell:
    value: object = None          # float for metrics, MARK / "" for markers
    text: str = ""
    key: Optional[str] = None     # the results-table column it came from
    text_color: Optional[str] = None
    background: Optional[str] = None
    bold: bool = False
    tooltip: Optional[str] = None


@dataclass
class ViewRow:
    label: str
    kind: str                     # "metric" | "marker"
    cells: list
    xl_format: Optional[str] = None


@dataclass
class QuarterView:
    symbol: str
    period: Optional[str]
    quarters: list = field(default_factory=list)
    rows: list = field(default_factory=list)
    note: Optional[str] = None

    @property
    def empty(self) -> bool:
        return not self.quarters or not self.rows


def _num(v):
    """A finite float, or None for missing / NaN / non-numeric."""
    if v is None:
        return None
    try:
        f = float(v)
    except (TypeError, ValueError):
        return None
    return None if math.isnan(f) else f


def _ts(v):
    if v is None:
        return None
    try:
        if pd.isna(v):
            return None
        return pd.Timestamp(v)
    except (TypeError, ValueError):
        return None


def _metric_cell(row: dict, key: str, fmt: str, styles) -> Cell:
    value = _num(row.get(key))
    text = "N/A" if value is None else fmt.format(value)
    tcol, bg, bold = (styles.resolve(key) if styles is not None
                      else (None, None, False))
    return Cell(value=value, text=text, key=key, text_color=tcol,
                background=bg, bold=bold, tooltip=cell_tooltip(key, row))


def _block_quarters(row: dict) -> list:
    """The Q-X quarters this ticker HAS in the scan: k with a report date on
    either side. A frame's columns run to the widest ticker, so a ticker with
    fewer reports carries empty blocks past its last one."""
    ks = sorted({int(m.group(1)) for key in row
                 if (m := _Q_COL_RE.match(str(key)))
                 and m.group(2) in ("report_date_eps", "report_date_rev")})
    out = []
    for k in ks:
        rd = _ts(row.get(f"q{k}_report_date_eps")) or _ts(
            row.get(f"q{k}_report_date_rev"))
        if rd is not None:
            out.append(QuarterHead(k, rd, _ts(row.get(f"_q{k}_period_ending"))))
    return out


def build_quarter_view(row: dict, *, styles=None, hidden_types=(),
                       hidden_keys=(), period: Optional[str] = None
                       ) -> QuarterView:
    """The Quarter view of one results row.

    `styles` is the table's `RowStyles` for the row (colours are reused, not
    re-evaluated, so they match the table exactly). `hidden_types` is the
    EFFECTIVE Hide Q Columns set; `hidden_keys` the individually hidden
    columns.
    """
    hidden_types = set(hidden_types or ())
    hidden_keys = set(hidden_keys or ())
    view = QuarterView(symbol=str(row.get("symbol", "")), period=period)
    quarters = _block_quarters(row)

    if quarters:
        view.quarters = quarters
        ks = [q.k for q in quarters]
        for side in ("eps", "rev"):
            for suffix, label, fmt, xl in METRICS[side]:
                keys = [f"q{k}_{suffix}" for k in ks]
                if not any(key in row for key in keys):
                    continue          # this side's blocks are not in the scan
                if f"q_{suffix}" in hidden_types:
                    continue
                if all(key in hidden_keys for key in keys):
                    continue
                view.rows.append(ViewRow(
                    label, "metric",
                    [_metric_cell(row, key, fmt, styles) for key in keys], xl))
        for prefix, run_label in C.RUN_SOURCES:
            qs = row.get(C.run_key(prefix))
            if isinstance(qs, np.ndarray):
                qs = qs.tolist()
            if not isinstance(qs, (list, tuple)):
                continue              # that filter did not run
            if f"series_{prefix}" in hidden_types:
                continue
            counted = {int(k) for k in list(qs)}
            view.rows.append(ViewRow(
                f"In {run_label}", "marker",
                [Cell(value=MARK if k in counted else "",
                      text=MARK if k in counted else "",
                      tooltip=(f"Q-{k} counted in the {run_label}"
                               if k in counted else None))
                 for k in ks]))
        if any(c.key in hidden_keys for r in view.rows for c in r.cells
               if c.key):
            view.note = ("Cells hidden in the results table are shown "
                         "without their colour.")
        if not view.rows:
            view.note = ("Every per-quarter metric of this scan is hidden "
                         "(Hide Q Columns).")
        return view

    # No Q-X blocks: the Current rows are the most recent quarter, if any.
    present = [s for s in _CURRENT_KEYS
               if s in row and s not in hidden_keys]
    if present:
        rd = _ts(row.get("last_report_date"))
        view.quarters = [QuarterHead(1, rd, _ts(row.get("_last_period_ending")))]
        for side in ("eps", "rev"):
            for suffix, label, fmt, xl in METRICS[side]:
                if suffix in present:
                    view.rows.append(ViewRow(
                        label, "metric", [_metric_cell(row, suffix, fmt, styles)],
                        xl))
        view.note = ("This scan has no quarter blocks, so only the most "
                     "recent quarter is shown. Turn on a Consecutive Beats or "
                     "YoY Growth filter to see more quarters.")
        return view

    view.note = ("No per-quarter earnings in this scan. Turn on a Consecutive "
                 "Beats or YoY Growth filter (display-only is enough) to see "
                 "this ticker's quarters.")
    return view


# ======================================================================
# Copy / Excel
# ======================================================================

def _header_rows(view: QuarterView) -> list:
    title = view.symbol + (f" · {view.period}" if view.period else "")
    return [
        [title] + [f"Q-{q.k}" for q in view.quarters],
        ["Reported"] + [q.report_date.strftime("%Y-%m-%d")
                        if q.report_date is not None else ""
                        for q in view.quarters],
        ["Fiscal quarter"] + [q.fiscal.strftime("%Y-%m")
                              if q.fiscal is not None else ""
                              for q in view.quarters],
    ]


def view_as_tsv(view: QuarterView) -> str:
    """Tab-separated text, as shown — pastes straight into Excel."""
    lines = ["\t".join(r) for r in _header_rows(view)]
    for row in view.rows:
        lines.append("\t".join([row.label] + [c.text for c in row.cells]))
    return "\n".join(lines) + "\n"


def write_quarter_view_xlsx(view: QuarterView, path, *,
                            colors: bool = True) -> None:
    """One sheet: three header rows (quarter, report date, fiscal quarter),
    then a row per metric with REAL numbers in the table's number formats,
    and the ✓ rows. With `colors`, each cell gets the table's text colour,
    fill and bold."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from .exports import ExportsController, _neutralize_formula

    wb = Workbook()
    ws = wb.active
    ws.title = ExportsController._sanitize_sheet_name(
        view.symbol or "Quarters", set())
    header = _header_rows(view)
    bold = Font(bold=True)
    ws.append([_neutralize_formula(header[0][0])] + header[0][1:])
    ws.append([header[1][0]] + [
        q.report_date.to_pydatetime().date() if q.report_date is not None
        else None for q in view.quarters])
    ws.append([header[2][0]] + header[2][1:])
    for r in range(1, 4):
        for c in range(1, len(view.quarters) + 2):
            ws.cell(row=r, column=c).font = bold
    for c in range(2, len(view.quarters) + 2):
        ws.cell(row=2, column=c).number_format = "yyyy-mm-dd"

    for row in view.rows:
        ws.append([row.label] + [c.value if c.value != "" else None
                                 for c in row.cells])
        r = ws.max_row
        ws.cell(row=r, column=1).font = bold
        for j, cell in enumerate(row.cells, start=2):
            xc = ws.cell(row=r, column=j)
            if row.kind == "marker":
                xc.alignment = Alignment(horizontal="center")
                continue
            if row.xl_format and cell.value is not None:
                xc.number_format = row.xl_format
            if not colors:
                continue
            if cell.text_color or cell.bold:
                xc.font = Font(
                    bold=cell.bold,
                    color=(f"FF{cell.text_color[1:].upper()}"
                           if cell.text_color else None))
            if cell.background:
                argb = f"FF{cell.background[1:].upper()}"
                xc.fill = PatternFill(fill_type="solid", start_color=argb,
                                      end_color=argb)
    ws.freeze_panes = "B4"
    ws.column_dimensions["A"].width = max(
        14, max((len(r.label) for r in view.rows), default=0) + 2)
    for c in range(2, len(view.quarters) + 2):
        ws.column_dimensions[get_column_letter(c)].width = 12
    wb.save(path)


def default_export_name(view: QuarterView, today: Optional[date] = None) -> str:
    from .exports import ExportsController
    safe = ExportsController._safe_filename_component
    stamp = (today or date.today()).isoformat()
    period = f"_{safe(view.period)}" if view.period else ""
    return f"{safe(view.symbol)}_quarters{period}_{stamp}.xlsx"


# ======================================================================
# The window
# ======================================================================

class QuarterViewDialog(QDialog):
    """Non-modal, one per double-click, deleted on close."""

    def __init__(self, view: QuarterView, parent=None,
                 on_exported: Optional[Callable] = None):
        super().__init__(parent)
        self.view = view
        self._on_exported = on_exported
        self.setAttribute(Qt.WidgetAttribute.WA_DeleteOnClose)
        self.setModal(False)
        self.setWindowTitle(" — ".join(
            p for p in (view.symbol, view.period, "Quarter view") if p))
        lay = QVBoxLayout(self)

        self.note = QLabel(view.note or "")
        self.note.setWordWrap(True)
        self.note.setStyleSheet("color: #999;")
        self.note.setVisible(bool(view.note))
        lay.addWidget(self.note)

        self.table = QTableWidget(len(view.rows), len(view.quarters))
        self.table.setEditTriggers(QTableWidget.EditTrigger.NoEditTriggers)
        self.table.setHorizontalHeaderLabels(
            ["\n".join(q.label_lines()) for q in view.quarters])
        self.table.setVerticalHeaderLabels([r.label for r in view.rows])
        for i, row in enumerate(view.rows):
            for j, cell in enumerate(row.cells):
                item = QTableWidgetItem(cell.text)
                item.setTextAlignment(
                    Qt.AlignmentFlag.AlignCenter if row.kind == "marker"
                    else Qt.AlignmentFlag.AlignRight
                    | Qt.AlignmentFlag.AlignVCenter)
                if cell.text_color:
                    item.setForeground(QColor(cell.text_color))
                if cell.background:
                    item.setBackground(QColor(cell.background))
                if cell.bold:
                    font = item.font()
                    font.setBold(True)
                    item.setFont(font)
                if cell.tooltip:
                    item.setToolTip(cell.tooltip)
                self.table.setItem(i, j, item)
        self.table.resizeColumnsToContents()
        self.table.resizeRowsToContents()
        self.table.setVisible(not view.empty)
        lay.addWidget(self.table, 1)

        buttons = QHBoxLayout()
        self.btn_copy = QPushButton("Copy")
        self.btn_copy.setToolTip("Copy the grid as tab-separated text "
                                 "(pastes into Excel).")
        self.btn_export = QPushButton("Export to Excel…")
        self.chk_colors = QCheckBox("Include colours")
        self.chk_colors.setChecked(True)
        self.btn_close = QPushButton("Close")
        for w in (self.btn_copy, self.btn_export, self.chk_colors):
            w.setEnabled(not view.empty)
            buttons.addWidget(w)
        buttons.addStretch()
        buttons.addWidget(self.btn_close)
        lay.addLayout(buttons)
        self.status = QLabel("")
        self.status.setStyleSheet("color: #999; font-size: 11px;")
        lay.addWidget(self.status)

        self.btn_copy.clicked.connect(self.copy_to_clipboard)
        self.btn_export.clicked.connect(self.export_dialog)
        self.btn_close.clicked.connect(self.close)
        self._fit_to_content()

    def _fit_to_content(self):
        """Size to the grid, never past 90% of the screen (no nativeEvent —
        it crashes PyQt 6.7.1; the main window's fit logic is not needed for
        a small dialog)."""
        t = self.table
        width = (t.verticalHeader().width() + t.frameWidth() * 2 + 40
                 + sum(t.columnWidth(c) for c in range(t.columnCount())))
        height = (t.horizontalHeader().height() + t.frameWidth() * 2 + 140
                  + sum(t.rowHeight(r) for r in range(t.rowCount())))
        screen = (self.parent().screen() if self.parent() is not None
                  else QApplication.primaryScreen())
        if screen is not None:
            avail = screen.availableGeometry()
            width = min(width, int(avail.width() * 0.9))
            height = min(height, int(avail.height() * 0.9))
        self.resize(max(width, 420), max(height, 200))

    def copy_to_clipboard(self):
        QApplication.clipboard().setText(view_as_tsv(self.view))
        self.status.setText("Copied — paste into Excel.")

    def export_dialog(self):
        path, _ = QFileDialog.getSaveFileName(
            self, "Export Quarter View", default_export_name(self.view),
            "Excel Workbook (*.xlsx)")
        if not path:
            return
        try:
            write_quarter_view_xlsx(self.view, path,
                                    colors=self.chk_colors.isChecked())
        except ImportError as exc:
            QMessageBox.warning(self, "Export Error",
                                f"XLSX export requires the openpyxl package.\n{exc}")
            return
        except OSError as exc:
            QMessageBox.warning(self, "Export Error",
                                f"Could not write file:\n{exc}")
            return
        self.status.setText(f"Exported to {path}")
        if self._on_exported is not None:
            try:
                self._on_exported(path, self.view)
            except Exception as exc:
                log.debug("quarter view export callback failed: %s", exc)
