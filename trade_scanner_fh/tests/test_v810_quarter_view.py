"""v8.1.0 — Quarter view: double-click a ticker to see its quarters as columns.

User decisions (2026-10-04) these tests pin:
  * a PLAIN left double-click opens it; one window per double-click; a
    snapshot (later scans / hides do not change an open view);
  * columns = every quarter the scan produced for the ticker, regardless of
    hiding; headers carry the report date and fiscal quarter (no Date rows);
  * rows = per-quarter metrics in the scan, dropped when the Hide Q type is
    ticked or EVERY quarter of the metric is hidden individually;
  * a ✓ row per active Beats / Growth / Accel filter under its counted
    quarters; the table's own colours; tooltips;
  * no Q-X blocks → Q-1 from the Current rows, or a note;
  * Copy (TSV) and Export to Excel (one sheet, real numbers, colours
    optional).
"""
from __future__ import annotations

import numpy as np
import openpyxl
import pandas as pd
import pytest
from PyQt6.QtCore import Qt
from PyQt6.QtTest import QTest

from trade_scanner_fh.gui import coloring as C
from trade_scanner_fh.gui import quarter_view as Q
from trade_scanner_fh.gui import widgets as W

T = pd.Timestamp
DATES = [("2026-05-05", "2026-03-01"), ("2026-02-10", "2025-12-01"),
         ("2025-11-04", "2025-09-01"), ("2025-08-05", "2025-06-01")]


def _row(n_q=4, symbol="AAA", rev=True, extra_empty=0, **extra):
    """A results-row dict shaped like the scanner's: Q-X blocks for EPS (and
    Rev), fiscal quarters, plus `extra_empty` blocks past the ticker's last
    report (a frame's columns run to the widest ticker)."""
    row = {"symbol": symbol}
    sides = ("eps", "rev") if rev else ("eps",)
    for k in range(1, n_q + extra_empty + 1):
        real = k <= n_q
        d, fp = DATES[(k - 1) % len(DATES)]
        for s in sides:
            row[f"q{k}_report_date_{s}"] = T(d) if real else pd.NaT
            row[f"q{k}_reported_{s}"] = float(k) if real else np.nan
            row[f"q{k}_surprise_{s}_dollar"] = 0.1 * k if real else np.nan
            row[f"q{k}_surprise_{s}_pct"] = 10.0 * k if real else np.nan
            row[f"q{k}_yoy_{s}_pct"] = 5.0 * k if real else np.nan
        row[f"_q{k}_period_ending"] = T(fp) if real else pd.NaT
    row.update(extra)
    return row


def _labels(view):
    return [r.label for r in view.rows]


# ======================================================================
# The layout rules (pure)
# ======================================================================

def test_quarters_are_the_tickers_own_with_dates_in_the_header():
    view = Q.build_quarter_view(_row(n_q=3, extra_empty=2), period="1D")
    assert [q.k for q in view.quarters] == [1, 2, 3], \
        "empty blocks past the ticker's last report are not quarters"
    assert view.quarters[1].label_lines() == ["Q-2", "2026-02-10",
                                              "FQ 2025-12"]
    assert "Q-X Date" not in " ".join(_labels(view))


def test_metric_rows_follow_the_sides_in_the_scan():
    assert _labels(Q.build_quarter_view(_row())) == [
        "Reported EPS", "Surp EPS $", "Surp EPS %", "YoY EPS %",
        "Reported Rev", "Surp Rev $", "Surp Rev %", "YoY Rev %"]
    assert _labels(Q.build_quarter_view(_row(rev=False))) == [
        "Reported EPS", "Surp EPS $", "Surp EPS %", "YoY EPS %"]


def test_values_read_like_the_table_and_missing_is_na():
    row = _row(n_q=2)
    row["q2_yoy_eps_pct"] = np.nan
    view = Q.build_quarter_view(row)
    by = {r.label: r for r in view.rows}
    assert [c.text for c in by["Surp EPS %"].cells] == ["+10.00%", "+20.00%"]
    assert [c.text for c in by["Reported Rev"].cells] == ["1.0", "2.0"]
    assert by["YoY EPS %"].cells[1].text == "N/A"
    assert by["YoY EPS %"].cells[1].value is None
    assert by["Surp EPS $"].cells[0].key == "q1_surprise_eps_dollar"


def test_a_hidden_type_drops_its_metric_row():
    view = Q.build_quarter_view(_row(), hidden_types={"q_surprise_eps_dollar",
                                                      "q_yoy_rev_pct"})
    assert "Surp EPS $" not in _labels(view)
    assert "YoY Rev %" not in _labels(view)
    assert "Surp Rev $" in _labels(view)


def test_individual_hides_drop_a_metric_only_when_every_quarter_is_hidden():
    keys = {f"q{k}_reported_eps" for k in range(1, 5)}
    assert "Reported EPS" not in _labels(
        Q.build_quarter_view(_row(), hidden_keys=keys))
    partial = Q.build_quarter_view(_row(), hidden_keys={"q2_reported_eps"})
    row = next(r for r in partial.rows if r.label == "Reported EPS")
    assert [c.text for c in row.cells] == ["1.00", "2.00", "3.00", "4.00"], \
        "quarters are shown regardless of hiding"
    assert partial.note and "without their colour" in partial.note


def test_marker_rows_tick_the_quarters_each_run_counted():
    view = Q.build_quarter_view(_row(
        _consec_eps_beats_qs=[1, 2], _accel_eps_surp_qs=np.array([2, 3]),
        _consec_rev_growth_qs=[]))
    marks = {r.label: [c.text for c in r.cells]
             for r in view.rows if r.kind == "marker"}
    assert marks == {
        "In EPS beats streak": ["✓", "✓", "", ""],
        "In YoY Rev growth run": ["", "", "", ""],
        "In Accel EPS surprise series": ["", "✓", "✓", ""],
    }
    hidden = Q.build_quarter_view(_row(_consec_eps_beats_qs=[1]),
                                  hidden_types={"series_consec_eps_beats"})
    assert not [r for r in hidden.rows if r.kind == "marker"]


def test_cells_take_the_tables_colours():
    rs = C.RowStyles()
    rs.set_cell("q2_surprise_eps_pct", "text", 3, "#4caf50")
    rs.set_cell("q2_surprise_eps_pct", "background", 2, "#1d3f5c")
    rs.set_cell("q2_surprise_eps_pct", "bold", 1, True)
    view = Q.build_quarter_view(_row(), styles=rs)
    cell = next(r for r in view.rows if r.label == "Surp EPS %").cells[1]
    assert (cell.text_color, cell.background, cell.bold) == \
        ("#4caf50", "#1d3f5c", True)
    other = next(r for r in view.rows if r.label == "Surp EPS %").cells[0]
    assert (other.text_color, other.background, other.bold) == \
        (None, None, False)


def test_cells_carry_the_hover_tooltip():
    view = Q.build_quarter_view(_row())
    cell = next(r for r in view.rows if r.label == "Reported EPS").cells[2]
    assert cell.tooltip == "Q-3 · reported 2025-11-04 · fiscal quarter 2025-09"


def test_without_blocks_the_current_rows_give_q1():
    row = {"symbol": "CUR", "reported_eps": 1.5, "surprise_rev_pct": 4.0,
           "last_report_date": T("2026-05-05"),
           "_last_period_ending": T("2026-03-01")}
    view = Q.build_quarter_view(row)
    assert [q.label_lines() for q in view.quarters] == [
        ["Q-1", "2026-05-05", "FQ 2026-03"]]
    assert _labels(view) == ["Reported EPS", "Surp Rev %"]
    assert view.rows[0].cells[0].text == "1.50"
    assert "no quarter blocks" in view.note


def test_nothing_to_show_says_how_to_get_quarters():
    view = Q.build_quarter_view({"symbol": "NIL", "close": 3.0})
    assert view.empty and "Consecutive Beats or YoY Growth" in view.note
    hidden_all = Q.build_quarter_view(_row(rev=False), hidden_types={
        "q_reported_eps", "q_surprise_eps_dollar", "q_surprise_eps_pct",
        "q_yoy_eps_pct"})
    assert hidden_all.empty and "hidden" in hidden_all.note


def test_copy_text_is_tab_separated_as_shown():
    view = Q.build_quarter_view(_row(n_q=2, rev=False,
                                     _consec_eps_beats_qs=[1]), period="1D")
    lines = Q.view_as_tsv(view).splitlines()
    assert lines[:3] == ["AAA · 1D\tQ-1\tQ-2",
                         "Reported\t2026-05-05\t2026-02-10",
                         "Fiscal quarter\t2026-03\t2025-12"]
    assert lines[5] == "Surp EPS %\t+10.00%\t+20.00%"
    assert lines[-1] == "In EPS beats streak\t✓\t"


def test_default_export_name_is_file_safe():
    view = Q.QuarterView(symbol="BRK/B", period="Custom 09/01 → 09/18")
    name = Q.default_export_name(view, today=pd.Timestamp("2026-10-04").date())
    assert name == "BRK_B_quarters_Custom 09_01 → 09_18_2026-10-04.xlsx"


# ======================================================================
# Excel
# ======================================================================

def _styled_view():
    rs = C.RowStyles()
    rs.set_cell("q1_surprise_eps_pct", "text", 2, "#4caf50")
    rs.set_cell("q1_surprise_eps_pct", "background", 1, "#1d3f5c")
    rs.set_cell("q1_surprise_eps_pct", "bold", 1, True)
    return Q.build_quarter_view(
        _row(n_q=3, rev=False, symbol="BRK/B", _consec_eps_beats_qs=[1, 2]),
        styles=rs, period="1D")


def test_excel_sheet_has_real_dates_numbers_formats_and_marks(tmp_path):
    path = tmp_path / "q.xlsx"
    Q.write_quarter_view_xlsx(_styled_view(), path)
    ws = openpyxl.load_workbook(path).active
    assert ws.title == "BRK_B"
    rows = list(ws.iter_rows(values_only=True))
    assert rows[0] == ("BRK/B · 1D", "Q-1", "Q-2", "Q-3")
    assert rows[1][1].date() == T("2026-05-05").date()
    assert ws.cell(row=2, column=2).number_format == "yyyy-mm-dd"
    assert rows[2] == ("Fiscal quarter", "2026-03", "2025-12", "2025-09")
    assert rows[5] == ("Surp EPS %", 10, 20, 30)
    assert ws.cell(row=6, column=2).number_format == Q._PCT_XL
    assert ws.cell(row=4, column=2).number_format == "0.00"
    assert rows[-1] == ("In EPS beats streak", "✓", "✓", None)
    assert ws.freeze_panes == "B4"


def test_excel_colours_follow_the_checkbox(tmp_path):
    on, off = tmp_path / "on.xlsx", tmp_path / "off.xlsx"
    Q.write_quarter_view_xlsx(_styled_view(), on, colors=True)
    Q.write_quarter_view_xlsx(_styled_view(), off, colors=False)
    c_on = openpyxl.load_workbook(on).active.cell(row=6, column=2)
    c_off = openpyxl.load_workbook(off).active.cell(row=6, column=2)
    assert c_on.font.color.rgb == "FF4CAF50" and c_on.font.bold
    assert c_on.fill.fill_type == "solid" and c_on.fill.start_color.rgb == "FF1D3F5C"
    assert c_off.fill.fill_type is None and not c_off.font.bold
    assert openpyxl.load_workbook(on).active.cell(row=6, column=3).fill.fill_type is None


# ======================================================================
# The table — which double-clicks open a view
# ======================================================================

def _frame():
    rows = [_row(n_q=3, symbol=s, rev=False) for s in ("AAA", "MMM", "ZZZ")]
    for i, r in enumerate(rows):
        r.update(close=10.0 + i, pct_gain=1.0, gain_start_date=T("2026-09-01"),
                 consec_eps_beats=2, _consec_eps_beats_qs=[1, 2])
    rows[2]["q1_reported_eps"] = 99.0
    return pd.DataFrame(rows)


@pytest.fixture
def table(_qapp):
    t = W.ResultsTable()
    t.resize(1600, 400)
    t.show()
    t.populate(_frame())
    t.sortByColumn(0, Qt.SortOrder.DescendingOrder)      # ZZZ on top
    yield t
    t.close()


def _dclick(t, view_row, button=Qt.MouseButton.LeftButton,
            mods=Qt.KeyboardModifier.NoModifier):
    seen = []
    t.quarter_view_requested.connect(seen.append)
    pos = t.visualRect(t.proxy.index(view_row, 0)).center()
    QTest.mouseDClick(t.viewport(), button, mods, pos)
    t.quarter_view_requested.disconnect(seen.append)
    return seen


def test_a_plain_double_click_names_the_frame_row_under_the_sort(table):
    assert table.proxy.index(0, 0).data() == "ZZZ"
    assert _dclick(table, 0) == [2], "ZZZ is frame row 2"


@pytest.mark.parametrize("button, mods", [
    (Qt.MouseButton.LeftButton, Qt.KeyboardModifier.ShiftModifier),
    (Qt.MouseButton.LeftButton, Qt.KeyboardModifier.ControlModifier),
    (Qt.MouseButton.RightButton, Qt.KeyboardModifier.NoModifier),
    (Qt.MouseButton.MiddleButton, Qt.KeyboardModifier.NoModifier),
])
def test_hotkey_style_double_clicks_do_not_open_a_view(table, button, mods):
    assert _dclick(table, 0, button, mods) == []


def test_nothing_opens_while_a_render_is_filling(table):
    table._populate_in_flight = True
    try:
        assert _dclick(table, 0) == []
    finally:
        table._populate_in_flight = False


def test_row_snapshot_is_a_copy_that_survives_a_new_render(table):
    row, styles = table.row_snapshot(2)
    assert row["symbol"] == "ZZZ" and row["q1_reported_eps"] == 99.0
    assert styles is not None, "default rules colour the beats quarters"
    table.populate(_frame().assign(q1_reported_eps=0.0))
    assert row["q1_reported_eps"] == 99.0
    assert table.row_snapshot(99) is None


# ======================================================================
# The window
# ======================================================================

@pytest.fixture
def window(_qapp, tmp_parquets, monkeypatch):
    from trade_scanner_fh.gui import main_window as mw_mod
    monkeypatch.setattr(mw_mod, "PRESETS_DIR", tmp_parquets / "presets")
    (tmp_parquets / "presets").mkdir(exist_ok=True)
    monkeypatch.setattr(mw_mod.scan_history, "record_scan_results",
                        lambda *a, **k: {})
    w = mw_mod.MainWindow()
    yield w
    for d in list(w._quarter_views):
        d.close()
    w.close()
    w.deleteLater()


def _scan(window, df):
    from trade_scanner_fh.gui.workers import WorkerScanResult
    window._on_scan_done(WorkerScanResult(period_results={"1D": df},
                                          period_order=["1D"]))


def _grid(dlg):
    t = dlg.table
    return {t.verticalHeaderItem(r).text():
            [t.item(r, c).text() for c in range(t.columnCount())]
            for r in range(t.rowCount())}


def test_each_double_click_opens_its_own_window(window):
    _scan(window, _frame())
    window.results_table.quarter_view_requested.emit(0)
    window.results_table.quarter_view_requested.emit(2)
    titles = [d.windowTitle() for d in window._quarter_views]
    assert titles == ["AAA — 1D — Quarter view", "ZZZ — 1D — Quarter view"]
    window._quarter_views[0].close()
    assert [d.windowTitle() for d in window._quarter_views] == \
        ["ZZZ — 1D — Quarter view"]


def test_the_view_paints_exactly_what_the_table_paints(window):
    _scan(window, _frame())
    dlg = window._open_quarter_view(0)
    t = window.results_table
    keys = [k for _h, k, _f in t.active_columns]
    src_col = keys.index("q1_surprise_eps_pct")
    table_fg = t.model_src.item(0, src_col).foreground().color().name()
    r = [dlg.table.verticalHeaderItem(i).text()
         for i in range(dlg.table.rowCount())].index("Surp EPS %")
    view_fg = dlg.table.item(r, 0).foreground().color().name()
    assert table_fg == view_fg == C.STREAK_GREEN, \
        "Q-1 is in the beats streak: default streak green, same in both"
    q3 = dlg.table.item(r, 2).foreground().color().name()
    assert q3 != C.STREAK_GREEN, "Q-3 is outside the streak"


def test_hides_shape_the_metrics_not_the_quarters(window):
    _scan(window, _frame())
    window._on_hide_type_toggled("q_surprise_eps_dollar", True)
    window._on_columns_hide_requested(["q2_reported_eps"])
    grid = _grid(window._open_quarter_view(0))
    assert "Surp EPS $" not in grid
    assert grid["Reported EPS"] == ["1.00", "2.00", "3.00"]


def test_an_open_view_is_a_snapshot(window):
    _scan(window, _frame())
    dlg = window._open_quarter_view(2)
    _scan(window, _frame().assign(q1_reported_eps=0.0))
    assert _grid(dlg)["Reported EPS"][0] == "99.00"


def test_copy_and_export_from_the_window(window, tmp_path, monkeypatch):
    _scan(window, _frame())
    dlg = window._open_quarter_view(0)
    dlg.btn_copy.click()
    from PyQt6.QtWidgets import QApplication
    assert QApplication.clipboard().text() == Q.view_as_tsv(dlg.view)

    target = tmp_path / "out.xlsx"
    monkeypatch.setattr(Q.QFileDialog, "getSaveFileName",
                        staticmethod(lambda *a, **k: (str(target), "")))
    lines = []
    window.log_panel.write_line = lines.append
    dlg.chk_colors.setChecked(False)
    dlg.btn_export.click()
    ws = openpyxl.load_workbook(target).active
    assert ws.cell(row=1, column=2).value == "Q-1"
    assert ws.cell(row=4, column=2).fill.fill_type is None, "colours off"
    assert any("Quarter view export: AAA" in ln for ln in lines)


def test_a_view_with_nothing_to_show_disables_its_buttons(window):
    df = pd.DataFrame([{"symbol": "NIL", "close": 1.0, "pct_gain": 1.0,
                        "gain_start_date": T("2026-09-01")}])
    _scan(window, df)
    dlg = window._open_quarter_view(0)
    assert not dlg.table.isVisibleTo(dlg) and dlg.note.isVisibleTo(dlg)
    assert not dlg.btn_export.isEnabled() and not dlg.btn_copy.isEnabled()
