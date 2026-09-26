"""v8.0.0 phase 3 — the Color Rules editor and its wiring into the window,
the table, presets and the Excel export."""
from __future__ import annotations

import json

import numpy as np
import pandas as pd
import pytest
from PyQt6.QtCore import Qt

from trade_scanner_fh.gui import coloring as C
from trade_scanner_fh.gui import color_rules_dialog as D
from trade_scanner_fh.gui import widgets as W


@pytest.fixture(scope="module", autouse=True)
def qapp():
    from PyQt6.QtWidgets import QApplication
    return QApplication.instance() or QApplication([])


COLS = [("RVOL", "rvol", "num"), ("% Gain", "pct_gain", "num"),
        ("Close", "close", "num"), ("Max Gap Date", "max_gap_date", "date"),
        ("FV Index", "index_membership", "text"),
        ("Consec EPS Beats", "consec_eps_beats", "num"),
        ("Q-X Surp EPS %", "q{k}_surprise_eps_pct", "num"),
        ("Q-X Date (EPS)", "q{k}_report_date_eps", "date")]
TARGETS = [("Q-X Reported EPS", "q_reported_eps"), ("RVOL", "rvol"),
           ("Close", "close"), ("% Gain", "pct_gain")]


def _dialog(rules, preview=None):
    return D.ColorRulesDialog(rules, COLS, TARGETS, preview_fn=preview)


def _complex_rules():
    return [
        C.Rule(id="r1", name="Leaders", match="all",
               conditions=[C.Condition("value", "rvol", ">=", 2.0),
                           C.Condition("value", "pct_gain", "between",
                                       10.0, 2.5e9)],
               target="row",
               style=C.Style(background=C.ColorSpec("fixed", "#1b5e20"),
                             bold=True)),
        C.Rule(id="r2", name="Close vs SMA", match="any",
               conditions=[C.Condition("value", "close", ">", other="rvol"),
                           C.Condition("value", "index_membership",
                                       "contains", "S&P"),
                           C.Condition("value", "rvol", "top_pct", 10.0),
                           C.Condition("value", "rvol", "blank")],
               target="columns", target_columns=["close", "q_reported_eps"],
               style=C.Style(text=C.ColorSpec("random",
                                              palette=["#123456",
                                                       "#abcdef"]))),
        C.Rule(id="r3", name="Gap on earnings",
               conditions=[C.Condition("date", "max_gap_date", "within",
                                       other=C.ANY_REPORT_DATE, days=2)],
               target="matched", expand_units=True,
               style=C.Style(text=C.ColorSpec("random"),
                             background=C.ColorSpec("fixed", "#222222"))),
        C.Rule(id="r4", name="Big surprise", scope="quarter", enabled=False,
               conditions=[C.Condition("value", "q{k}_surprise_eps_pct",
                                       ">=", 50.0),
                           C.Condition("quarter", op="<=",
                                       other="consec_eps_beats"),
                           C.Condition("quarter", op=">", value=1.0)],
               target="columns", target_columns=["q_reported_eps"],
               style=C.Style(text=C.ColorSpec("fixed", "#ffcc00"))),
        C.Rule(id="r5", name="Fails", conditions=[
               C.Condition("filter", "rvol", "fails"),
               C.Condition("filter", C.ANY_FILTER, "passes")],
               match="any",
               style=C.Style(text=C.ColorSpec("fixed", "#ff0000"))),
        C.Rule(id="r6", name="Unknown column survives",
               conditions=[C.Condition("value", "hv_rank", ">", 1234567.89)],
               target="columns", target_columns=["not_in_scan"],
               style=C.Style(text=C.ColorSpec("fixed", "#00ff00"))),
    ]


# ======================================================================
# Round trip — opening the editor and pressing OK never changes a rule
# ======================================================================

@pytest.mark.parametrize("rules", [C.default_rules(), _complex_rules()],
                         ids=["defaults", "complex"])
def test_editor_round_trip_is_lossless(rules):
    dlg = _dialog(rules)
    out = dlg.rules()
    assert [r.to_dict() for r in out] == [r.to_dict() for r in rules]


def test_round_trip_survives_visiting_every_rule():
    rules = _complex_rules()
    dlg = _dialog(rules)
    for i in range(len(rules)):
        dlg.list.setCurrentRow(i)
    dlg.list.setCurrentRow(0)
    assert [r.to_dict() for r in dlg.rules()] == \
        [r.to_dict() for r in rules]


# ======================================================================
# Editing
# ======================================================================

def test_build_a_rule_through_the_widgets():
    dlg = _dialog([])
    dlg.add_rule()
    ed = dlg.editor
    ed.name.setText("RVOL burst")
    cond = ed.conditions[0]
    D._set(cond.kind, "value")
    D._set(cond.column, "rvol")
    D._set(cond.op, ">=")
    cond.value.setText("2.5")
    D._set(ed.target, "row")
    D._set(ed.background.mode, "fixed")
    ed.background.set_color("#0d47a1")
    ed.bold.setChecked(True)
    [rule] = dlg.rules()
    assert rule.name == "RVOL burst"
    assert rule.conditions[0].to_dict() == C.Condition(
        "value", "rvol", ">=", 2.5).to_dict()
    assert rule.target == "row"
    assert rule.style.background.mode == "fixed"
    assert rule.style.background.color == "#0d47a1"
    assert rule.style.bold is True


def test_numbers_accept_suffixes_and_columns_on_the_right():
    ed = D.ConditionEditor(COLS)
    D._set(ed.column, "close")
    D._set(ed.op, ">")
    ed.value.setText("2.5B")
    assert ed.condition().value == 2.5e9
    D._set(ed.rhs, "column")
    D._set(ed.other, "rvol")
    c = ed.condition()
    assert c.other == "rvol" and c.value is None


def test_quarter_scope_offers_the_q_templates_row_scope_does_not():
    ed = D.ConditionEditor(COLS)
    keys = [ed.column.itemData(i) for i in range(ed.column.count())]
    assert "q{k}_surprise_eps_pct" not in keys
    ed.set_scope("quarter")
    keys = [ed.column.itemData(i) for i in range(ed.column.count())]
    assert "q{k}_surprise_eps_pct" in keys


def test_date_condition_offers_any_report_date_and_only_dates():
    ed = D.ConditionEditor(COLS)
    D._set(ed.kind, "date")
    left = [ed.column.itemData(i) for i in range(ed.column.count())]
    right = [ed.other.itemData(i) for i in range(ed.other.count())]
    assert left == ["max_gap_date"]
    assert right[0] == C.ANY_REPORT_DATE


def test_list_buttons_add_duplicate_move_delete_restore():
    dlg = _dialog(_complex_rules()[:2])
    dlg.list.setCurrentRow(0)
    dlg.duplicate_rule()
    names = [r.name for r in dlg.rules()]
    assert names == ["Leaders", "Leaders (copy)", "Close vs SMA"]
    ids = [r.id for r in dlg.rules()]
    assert len(set(ids)) == 3, "a duplicate gets its own id"
    dlg.list.setCurrentRow(2)
    dlg.move_rule(-1)
    assert [r.name for r in dlg.rules()][1] == "Close vs SMA"
    dlg.delete_rule()
    assert len(dlg.rules()) == 2
    dlg.restore_defaults()
    assert [r.id for r in dlg.rules()] == [r.id for r in C.default_rules()]


def test_list_swatch_follows_style_edits():
    dlg = _dialog([])
    dlg.add_rule()                       # new rules start with yellow text
    ed = dlg.editor
    D._set(ed.background.mode, "fixed")
    ed.background.set_color("#1b3a57")
    icon = dlg.list.item(0).icon().pixmap(14, 14).toImage()
    assert icon.pixelColor(7, 7).name() == "#1b3a57"


def test_enable_checkbox_in_the_list():
    dlg = _dialog(C.default_rules())
    dlg.list.item(1).setCheckState(Qt.CheckState.Unchecked)
    assert [r.enabled for r in dlg.rules()] == \
        [True, False, True, True, True, True]


def test_palette_editor_add_remove_reset():
    pd_ = D.PaletteDialog(["#111111"])
    pd_.add_color("#222222")
    assert pd_.palette() == ["#111111", "#222222"]
    pd_.list.setCurrentRow(0)
    pd_._remove()
    assert pd_.palette() == ["#222222"]
    pd_._remove()
    assert pd_.palette() == ["#222222"], "never empties the palette"
    pd_._reset()
    assert pd_.palette() == list(C.DEFAULT_PALETTE)


def test_status_reports_rows_and_missing_columns():
    df = pd.DataFrame({"symbol": ["A", "B"], "rvol": [3.0, 1.0]})

    def preview(rules):
        report = {}
        C.evaluate(df, rules, ["symbol", "rvol"], report=report)
        return report
    rules = [C.Rule(id="a", name="ok",
                    conditions=[C.Condition("value", "rvol", ">=", 2.0)],
                    style=C.Style(text=C.ColorSpec("fixed", "#ff0000"))),
             C.Rule(id="b", name="needs",
                    conditions=[C.Condition("value", "hv_rank", ">", 1.0)],
                    style=C.Style(text=C.ColorSpec("fixed", "#ff0000")))]
    dlg = _dialog(rules, preview)
    dlg._refresh_status()
    assert dlg.list.item(0).text().endswith("1 rows")
    assert "needs hv_rank" in dlg.list.item(1).text()
    dlg.list.setCurrentRow(1)
    dlg._refresh_status()
    assert "hv_rank" in dlg.status.text()


def test_apply_emits_the_edited_rules():
    dlg = _dialog(C.default_rules())
    got = []
    dlg.rules_applied.connect(got.append)
    dlg.list.item(0).setCheckState(Qt.CheckState.Unchecked)
    dlg.apply()
    assert got and got[0][0].enabled is False


# ======================================================================
# Window, table, presets, export
# ======================================================================

@pytest.fixture
def window(tmp_parquets, monkeypatch):
    from trade_scanner_fh.gui import main_window as mw_mod
    monkeypatch.setattr(mw_mod, "PRESETS_DIR", tmp_parquets / "presets")
    (tmp_parquets / "presets").mkdir(exist_ok=True)
    w = mw_mod.MainWindow()
    yield w
    w.close()
    w.deleteLater()


def _frame():
    return pd.DataFrame([
        {"symbol": s, "close": c, "pct_gain": g, "rvol": r,
         "gain_start_date": pd.Timestamp("2026-01-02")}
        for s, c, g, r in (("AAA", 10.0, 20.0, 3.0), ("BBB", 5.0, 2.0, 1.0),
                           ("CCC", 7.0, 9.0, 2.5))])


def _load(w, periods):
    w._period_results = dict(periods)
    w._period_order = list(periods)
    w._active_period = w._period_order[0]
    w.combo_timeframe.blockSignals(True)
    w.combo_timeframe.clear()
    for label in periods:
        w.combo_timeframe.addItem(label, userData=label)
    w.combo_timeframe.blockSignals(False)
    w._on_timeframe_changed(0)


def _item(w, sym, key):
    t = w.results_table
    keys = [k for _h, k, _f in t.active_columns]
    c, sc = keys.index(key), keys.index("symbol")
    for r in range(t.model_src.rowCount()):
        if t.model_src.item(r, sc).text() == sym:
            return t.model_src.item(r, c)
    raise AssertionError(sym)


def _leaders():
    return C.Rule(id="lead", name="Leaders",
                  conditions=[C.Condition("value", "rvol", ">=", 2.0)],
                  target="row",
                  style=C.Style(background=C.ColorSpec("fixed", "#1b5e20"),
                                text=C.ColorSpec("fixed", "#ffffff"),
                                bold=True))


def test_applying_rules_repaints_background_text_and_bold(window):
    _load(window, {"1D": _frame()})
    window._apply_color_rules([_leaders()])
    cell = _item(window, "AAA", "close")
    assert cell.background().color().name() == "#1b5e20"
    assert cell.foreground().color().name() == "#ffffff"
    assert cell.font().bold() is True
    other = _item(window, "BBB", "close")
    assert other.font().bold() is False


def test_a_hidden_column_can_still_drive_a_rule(window):
    """Hiding removes a column from the LAYOUT only, so a rule can colour by
    a value the user chose not to look at."""
    _load(window, {"1D": _frame()})
    window._on_columns_hide_requested(["rvol"])
    window._apply_color_rules([_leaders()])
    assert _item(window, "AAA", "close").background().color().name() \
        == "#1b5e20"


def test_rules_are_saved_and_restored_per_preset(window, tmp_parquets):
    window._apply_color_rules([_leaders()], rerender=False)
    window.preset_combo.addItem("colors")
    window.preset_combo.setCurrentText("colors")
    window._save_preset()
    data = json.loads((tmp_parquets / "presets" / "colors.json")
                      .read_text(encoding="utf-8"))
    assert data["color_rules"]["rules"][0]["id"] == "lead"
    window._apply_color_rules(C.default_rules(), rerender=False)
    window._load_preset()
    assert [r.id for r in window._color_rules] == ["lead"]
    assert [r.id for r in window.results_table.color_rules] == ["lead"]


def test_a_preset_without_rules_gets_the_default_schemes(window, tmp_parquets):
    (tmp_parquets / "presets" / "old.json").write_text(json.dumps({
        "_preset_version": 6, "indicators": {}}), encoding="utf-8")
    window._apply_color_rules([_leaders()], rerender=False)
    window.preset_combo.addItem("old")
    window.preset_combo.setCurrentText("old")
    window._load_preset()
    assert [r.id for r in window._color_rules] == \
        [r.id for r in C.default_rules()]


def test_open_dialog_follows_a_preset_load(window, tmp_parquets):
    window._open_color_rules_dialog()
    dlg = window._color_rules_dialog
    (tmp_parquets / "presets" / "p.json").write_text(json.dumps({
        "_preset_version": 7, "indicators": {},
        "color_rules": C.rules_to_json([_leaders()])}), encoding="utf-8")
    window.preset_combo.addItem("p")
    window.preset_combo.setCurrentText("p")
    window._load_preset()
    assert [r.id for r in dlg.rules()] == ["lead"]
    dlg.close()


def test_dialog_apply_repaints_the_table(window):
    _load(window, {"1D": _frame()})
    window._open_color_rules_dialog()
    dlg = window._color_rules_dialog
    dlg.set_rules([_leaders()])
    dlg.apply()
    assert _item(window, "CCC", "close").background().color().name() \
        == "#1b5e20"
    dlg.close()


def test_column_choices_include_hidden_and_templates(window):
    _load(window, {"1D": _frame()})
    window._on_columns_hide_requested(["rvol"])
    keys = {k for _l, k, _kind in window._color_rule_column_choices()}
    assert {"rvol", "close", "q{k}_surprise_eps_pct"} <= keys
    targets = {k for _l, k in window._color_rule_target_choices()}
    assert "q_reported_eps" in targets and "q1_reported_eps" not in targets


def test_excel_colours_every_period_with_fill_and_bold(window, tmp_path):
    from openpyxl import load_workbook
    _load(window, {"1D": _frame(), "1W": _frame()})
    window._apply_color_rules([_leaders()])
    keys = [k for _h, k, _f in window._ordered_active_columns_for_export()]
    path = tmp_path / "c.xlsx"
    window._write_xlsx_multi_sheet(str(path), ["1D", "1W"], keys, False,
                                   apply_colors=True)
    wb = load_workbook(path)
    for ws in wb.worksheets:
        header = [c.value for c in ws[1]]
        close_col = header.index("Close") + 1
        rows = {ws.cell(row=r, column=1).value: ws.cell(row=r, column=close_col)
                for r in range(2, ws.max_row + 1)}
        assert rows["AAA"].fill.start_color.rgb == "FF1B5E20", ws.title
        assert rows["AAA"].font.bold is True
        assert rows["BBB"].fill.fill_type is None


def test_excel_colours_follow_a_column_drag(window, tmp_path):
    """v8.0.0 fix: colours were mapped in the layout's canonical order while
    values are written in the user's drag order."""
    from openpyxl import load_workbook
    _load(window, {"1D": _frame()})
    rule = C.Rule(id="rv", name="rvol only",
                  conditions=[C.Condition("value", "rvol", ">=", 2.0)],
                  target="columns", target_columns=["rvol"],
                  style=C.Style(background=C.ColorSpec("fixed", "#abcdef")))
    window._apply_color_rules([rule])
    keys = [k for _h, k, _f in window._ordered_active_columns_for_export()]
    dragged = ["rvol"] + [k for k in keys if k != "rvol"]
    path = tmp_path / "d.xlsx"
    window._write_xlsx_multi_sheet(str(path), ["1D"], dragged, False,
                                   apply_colors=True)
    ws = load_workbook(path).active
    header = [c.value for c in ws[1]]
    assert header[0] == "RVOL"
    row_aaa = next(r for r in range(2, ws.max_row + 1)
                   if ws.cell(row=r, column=2).value == "AAA")
    assert ws.cell(row=row_aaa, column=1).fill.start_color.rgb == "FFABCDEF"
    assert ws.cell(row=row_aaa, column=2).fill.fill_type is None


# ----------------------------------------------------------------------
# Export header collisions (found by the 2026-09-26 round-2 smoke test)
# ----------------------------------------------------------------------

def _both_sides_frame():
    rows = []
    for sym, rv in (("AAA", 3.0), ("BBB", 1.0)):
        row = {"symbol": sym, "close": 10.0, "pct_gain": 5.0, "rvol": rv,
               "gain_start_date": pd.Timestamp("2026-01-02"),
               "consec_eps_beats": 2, "consec_rev_beats": 2,
               "_consec_eps_beats_qs": [1, 2],
               "_consec_rev_beats_qs": [1, 2]}
        for k in (1, 2):
            # EPS and Rev dates differ on purpose, so an overwrite shows.
            row[f"q{k}_report_date_eps"] = pd.Timestamp(f"2026-0{k}-10")
            row[f"q{k}_report_date_rev"] = pd.Timestamp(f"2026-0{k}-20")
            for side in ("eps", "rev"):
                row[f"q{k}_reported_{side}"] = 1.0 + k
                row[f"q{k}_surprise_{side}_dollar"] = 0.1
                row[f"q{k}_surprise_{side}_pct"] = 5.0
                row[f"q{k}_yoy_{side}_pct"] = 10.0 * k
        rows.append(row)
    return pd.DataFrame(rows)


def test_unique_export_headers_rename_only_the_clashes():
    from trade_scanner_fh.gui.exports import _unique_export_headers
    got = _unique_export_headers([
        ("Ticker", "symbol"), ("Q-1 Date", "q1_report_date_eps"),
        ("Q-1 Date", "q1_report_date_rev"), ("X", "a"), ("X", "b")])
    assert [h for h, _k in got] == [
        "Ticker", "Q-1 Date (EPS)", "Q-1 Date (Rev)", "X [a]", "X [b]"]


def test_export_keeps_both_date_columns_with_their_own_values(window,
                                                              tmp_path):
    """Before: the frame was keyed by header, so the Rev "Q-k Date" overwrote
    the EPS one and a column vanished per quarter (7.0.2 and earlier too)."""
    from openpyxl import load_workbook
    _load(window, {"1D": _both_sides_frame()})
    keys = [k for _h, k, _f in window._ordered_active_columns_for_export()]
    path = tmp_path / "both.xlsx"
    window._write_xlsx_multi_sheet(str(path), ["1D"], keys, False)
    ws = load_workbook(path).active
    header = [c.value for c in ws[1]]
    assert len(header) == len(keys)
    assert "Q-1 Date (EPS)" in header and "Q-1 Date (Rev)" in header
    eps = ws.cell(row=2, column=header.index("Q-1 Date (EPS)") + 1).value
    rev = ws.cell(row=2, column=header.index("Q-1 Date (Rev)") + 1).value
    assert str(eps).startswith("2026-01-10") and str(rev).startswith("2026-01-20")
    csv_path = tmp_path / "both.csv"
    window._write_csv_export(str(csv_path), ["1D"], keys, True)
    head = pd.read_csv(csv_path).columns.tolist()
    assert "News_Q-1 Date (EPS)" in head and "News_Q-1 Date (Rev)" in head


def test_excel_colours_stay_on_their_cells_past_a_duplicate_header(window,
                                                                   tmp_path):
    from openpyxl import load_workbook
    _load(window, {"1D": _both_sides_frame()})
    rule = C.Rule(id="yr", name="rev yoy", scope="quarter",
                  conditions=[C.Condition(kind="quarter", op="in_run",
                                          other="consec_rev_beats")],
                  target="columns", target_columns=["q_yoy_rev_pct"],
                  style=C.Style(background=C.ColorSpec("fixed", "#123456")))
    window._apply_color_rules([rule])
    keys = [k for _h, k, _f in window._ordered_active_columns_for_export()]
    path = tmp_path / "c.xlsx"
    window._write_xlsx_multi_sheet(str(path), ["1D"], keys, False,
                                   apply_colors=True)
    ws = load_workbook(path).active
    header = [c.value for c in ws[1]]
    painted = {header[c.column - 1] for row in ws.iter_rows(min_row=2)
               for c in row if c.fill.fill_type == "solid"}
    assert painted == {"Q-1 YoY Rev %", "Q-2 YoY Rev %"}


def test_colour_column_list_matches_the_written_sheet_header_for_header(
        window):
    """The colour pass places by position against this list; it must be the
    exact column list the sheet is written with, News columns included."""
    _load(window, {"1D": _both_sides_frame()})
    keys = [k for _h, k, _f in window._ordered_active_columns_for_export()]
    df = window._period_results["1D"]
    for news in (False, True):
        listed = [h for h, _k, _n in
                  window._ordered_export_columns_with_news(keys, news)]
        written = list(window._build_export_df(df, keys, news).columns)
        assert listed == written
